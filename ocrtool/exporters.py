"""Simpan hasil OCR ke berbagai format file."""

from __future__ import annotations

import io
import json
import math
import os
import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from .pipeline import PageResult

# kode: (label, ekstensi, mimetype)
FORMATS: dict[str, tuple[str, str, str]] = {
    "txt": ("Teks (.txt)", ".txt", "text/plain; charset=utf-8"),
    "docx": ("Word (.docx)", ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": ("PDF searchable (gambar + teks)", ".pdf", "application/pdf"),
    "pdf_text": ("PDF teks saja", ".pdf", "application/pdf"),
    "md": ("Markdown (.md)", ".md", "text/markdown; charset=utf-8"),
    "json": ("JSON (teks + posisi kotak)", ".json", "application/json"),
}


def build_export(
    pages: list[PageResult],
    fmt: str,
    keep_layout: bool = False,
    combine: bool = True,
    texts: list[str | None] | None = None,
) -> tuple[str, bytes, str]:
    """Kembalikan (nama_file, isi, mimetype). `texts` berisi teks hasil edit pengguna (opsional)."""
    if fmt not in FORMATS:
        raise ValueError(f"Format tidak dikenal: {fmt}")
    if not pages:
        raise ValueError("Tidak ada halaman untuk disimpan.")
    texts = texts or [None] * len(pages)
    texts = [t if t is not None else p.get_text(keep_layout) for p, t in zip(pages, texts)]
    _, ext, mime = FORMATS[fmt]
    writer = _WRITERS[fmt]

    if combine or len(pages) == 1:
        stem = Path(pages[0].name).stem if len(pages) == 1 else "hasil_ocr"
        return stem + ext, writer(pages, texts, keep_layout), mime

    buf = io.BytesIO()
    used: set[str] = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for page, text in zip(pages, texts):
            zf.writestr(_unique_name(Path(page.name).stem, ext, used), writer([page], [text], keep_layout))
    return "hasil_ocr.zip", buf.getvalue(), "application/zip"


def _unique_name(stem: str, ext: str, used: set[str]) -> str:
    name, n = stem + ext, 2
    while name.lower() in used:
        name, n = f"{stem}_{n}{ext}", n + 1
    used.add(name.lower())
    return name


# ---------------------------------------------------------------- teks


def _write_txt(pages, texts, keep_layout) -> bytes:
    if len(pages) == 1:
        return texts[0].encode("utf-8")
    parts = [f"===== {p.name} =====\n{t}" for p, t in zip(pages, texts)]
    return "\n\n".join(parts).encode("utf-8")


def _write_md(pages, texts, keep_layout) -> bytes:
    parts = []
    for page, text in zip(pages, texts):
        body = f"```text\n{text}\n```" if keep_layout else text
        parts.append(f"## {page.name}\n\n{body}" if len(pages) > 1 else body)
    return ("\n\n".join(parts) + "\n").encode("utf-8")


def _write_json(pages, texts, keep_layout) -> bytes:
    data = [
        {
            "name": p.name,
            "width": p.width,
            "height": p.height,
            "text": t,
            "score": round(p.score, 4),
            "boxes": [b.to_dict() for b in p.boxes],
        }
        for p, t in zip(pages, texts)
    ]
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


# ---------------------------------------------------------------- Word


def _write_docx(pages, texts, keep_layout) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    if keep_layout:
        for section in doc.sections:
            section.left_margin = section.right_margin = Cm(1.5)

    for index, (page, text) in enumerate(zip(pages, texts)):
        if index:
            doc.add_page_break()
        if len(pages) > 1:
            doc.add_heading(page.name, level=2)

        if keep_layout:
            lines = text.split("\n")
            # Perkecil huruf agar baris terpanjang muat di lebar halaman (Consolas ~0.55 em per huruf)
            longest = max((len(line) for line in lines), default=1)
            size = max(6.0, min(10.0, (18 * 28.35) / (0.55 * max(longest, 1))))
            for line in lines:
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                run = para.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(round(size, 1))
        else:
            for block in re.split(r"\n\s*\n", text.strip()):
                para = doc.add_paragraph()
                for i, line in enumerate(block.split("\n")):
                    run = para.add_run(line)
                    if i < block.count("\n"):
                        run.add_break()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

_FONT_CACHE: dict[bool, str] = {}


def _font(mono: bool = False) -> str:
    """Daftarkan font TTF yang mendukung Unicode; jatuh ke font bawaan PDF kalau tidak ada."""
    if mono in _FONT_CACHE:
        return _FONT_CACHE[mono]
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    win_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    candidates = (
        [win_fonts / "consola.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
         "/usr/share/fonts/dejavu/DejaVuSansMono.ttf", "/usr/share/fonts/TTF/DejaVuSansMono.ttf"]
        if mono
        else [win_fonts / "arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
              "/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/TTF/DejaVuSans.ttf",
              "/Library/Fonts/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf"]
    )
    name = "Courier" if mono else "Helvetica"
    for path in candidates:
        if Path(path).is_file():
            try:
                pdfmetrics.registerFont(TTFont("OCRMono" if mono else "OCRSans", str(path)))
                name = "OCRMono" if mono else "OCRSans"
                break
            except Exception:
                continue
    _FONT_CACHE[mono] = name
    return name


def _safe(text: str, font: str) -> str:
    # Font bawaan PDF (Helvetica/Courier) hanya mendukung Latin-1
    if font in ("Helvetica", "Courier"):
        return text.encode("latin-1", "replace").decode("latin-1")
    return text


def _write_pdf_searchable(pages, texts, keep_layout) -> bytes:
    """Gambar asli tampil di halaman, teks OCR diletakkan tak terlihat di atasnya
    sehingga bisa dicari (Ctrl+F) dan disalin. Teks hasil edit tidak dipakai di sini."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    font = _font()
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf)
    pdf.setTitle("Hasil OCR")
    for page in pages:
        a4_w, a4_h = A4 if page.height >= page.width else (A4[1], A4[0])
        scale = min(a4_w / page.width, a4_h / page.height)
        page_w, page_h = page.width * scale, page.height * scale
        pdf.setPageSize((page_w, page_h))
        pdf.drawImage(ImageReader(io.BytesIO(page.image_jpeg)), 0, 0, width=page_w, height=page_h)

        for box in page.boxes:
            text = _safe(box.text, font)
            (x0, y0), (x1, y1), _, (x3, y3) = box.box
            box_w = math.hypot(x1 - x0, y1 - y0) * scale
            box_h = math.hypot(x3 - x0, y3 - y0) * scale
            size = max(1.0, box_h * 0.8)
            text_w = stringWidth(text, font, size)
            if text_w <= 0 or box_w <= 0:
                continue
            pdf.saveState()
            pdf.translate(x3 * scale, page_h - y3 * scale)  # kiri-bawah kotak (sumbu y PDF naik ke atas)
            pdf.rotate(-math.degrees(math.atan2(y1 - y0, x1 - x0)))
            obj = pdf.beginText()
            obj.setTextRenderMode(3)  # tak terlihat, tapi bisa dipilih & dicari
            obj.setFont(font, size)
            obj.setHorizScale(100 * box_w / text_w)
            obj.setTextOrigin(0, box_h * 0.2)
            obj.textOut(text)
            pdf.drawText(obj)
            pdf.restoreState()
        pdf.showPage()
    pdf.save()
    return buf.getvalue()


def _write_pdf_text(pages, texts, keep_layout) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer

    font, mono = _font(), _font(mono=True)
    margin = 2 * cm if not keep_layout else 1.5 * cm
    avail_w = A4[0] - 2 * margin
    body = ParagraphStyle("body", fontName=font, fontSize=11, leading=15)
    heading = ParagraphStyle("heading", fontName=font, fontSize=13, leading=17, spaceAfter=8)

    story = []
    for index, (page, text) in enumerate(zip(pages, texts)):
        if index:
            story.append(PageBreak())
        if len(pages) > 1:
            story.append(Paragraph(escape(_safe(page.name, font)), heading))
        if keep_layout:
            text = _safe(text, mono)
            longest = max((len(line) for line in text.split("\n")), default=1)
            size = max(5.0, min(10.0, avail_w / max(1e-6, stringWidth("0" * max(longest, 1), mono, 1))))
            style = ParagraphStyle("mono", fontName=mono, fontSize=size, leading=size * 1.25)
            story.append(Preformatted(text, style))
        else:
            for block in re.split(r"\n\s*\n", _safe(text, font).strip()):
                story.append(Paragraph(escape(block).replace("\n", "<br/>"), body))
                story.append(Spacer(1, 6))
    if not story:
        story.append(Spacer(1, 1))

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
        title="Hasil OCR",
    ).build(story)
    return buf.getvalue()


_WRITERS = {
    "txt": _write_txt,
    "md": _write_md,
    "json": _write_json,
    "docx": _write_docx,
    "pdf": _write_pdf_searchable,
    "pdf_text": _write_pdf_text,
}
